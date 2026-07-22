"""Generate Diagram_Revision_Brief.docx from the markdown brief.

Produces a Word document version of diagram_revision_brief.md so groupmates
who prefer Word (or need to leave inline comments) can work from the same
spec. Runs standalone — no arguments.

Run:
    python "thesis_REMETHS/ThesisTemplate/_build_diagram_revision_brief.py"

Output: thesis_REMETHS/ThesisTemplate/Diagram_Revision_Brief.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Inches


HERE = Path(__file__).parent
OUT  = HERE / "Diagram_Revision_Brief.docx"

BRAND = (0x1F, 0x36, 0x5D)   # deep navy for headings
MUTED = (0x55, 0x5B, 0x67)
ACCENT_BG = "FFF3D6"          # pale amber for callouts
CODE_BG = "F1F3F5"


# ── Formatting helpers ─────────────────────────────────────────────────────
def _set_font(run, name: str = "Calibri", size: int = 11,
              bold: bool = False, italic: bool = False,
              color: tuple[int, int, int] | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _heading(doc: Document, text: str, level: int) -> None:
    sizes = {1: 20, 2: 15, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    _set_font(run, size=sizes.get(level, 11), bold=True, color=BRAND)


def _para(doc: Document, text: str, italic: bool = False,
          space_after: int = 6, size: int = 11) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    _set_font(run, italic=italic, size=size)


def _rich_para(doc: Document, segments: list[tuple[str, dict]],
               space_after: int = 6) -> None:
    """Segments = list of (text, kwargs_for_set_font)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    for text, kwargs in segments:
        run = p.add_run(text)
        _set_font(run, **kwargs)


def _bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75 + 0.75 * level)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    _set_font(run)


def _rich_bullet(doc: Document, segments: list[tuple[str, dict]],
                 level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75 + 0.75 * level)
    p.paragraph_format.space_after = Pt(2)
    for text, kwargs in segments:
        run = p.add_run(text)
        _set_font(run, **kwargs)


def _callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    _shade(cell, ACCENT_BG)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    _set_font(run, size=10, italic=True)
    doc.add_paragraph()   # breathing room after callout


def _code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.4)
    run = p.add_run(text)
    _set_font(run, name="Consolas", size=10, color=(0x2B, 0x2D, 0x42))


def _table(doc: Document, header: list[str], rows: list[list[str]],
           col_widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Light Grid Accent 1"
    for j, h in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        _shade(cell, "1F365D")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        _set_font(run, size=10, bold=True, color=(0xFF, 0xFF, 0xFF))
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            _set_font(run, size=10)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    if col_widths:
        for j, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[j].width = Inches(w)


def _hrule(doc: Document) -> None:
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B0B7C3")
    pBdr.append(bottom)
    pPr.append(pBdr)


# ── Document body ──────────────────────────────────────────────────────────
def build() -> None:
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    # Default body font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ── Title ──
    _heading(doc, "Diagram Revision Brief — Post-Defense Panel Comments", 1)
    _rich_para(doc, [
        ("Author: ", {"size": 10, "color": MUTED}),
        ("Sam", {"size": 10, "bold": True}),
        ("   ·   Date: ", {"size": 10, "color": MUTED}),
        ("2026-07-15", {"size": 10, "bold": True}),
        ("   ·   For: ", {"size": 10, "color": MUTED}),
        ("thesis groupmates", {"size": 10, "bold": True}),
    ])

    _para(doc,
        "The panel raised two diagram-specific revisions (comments #1 and #3 on "
        "the revisions form). Both are re-drawings. LaTeX prose and figure blocks "
        "are already in place in methodology.tex pointing at placeholder images; "
        "once you produce the finished PNGs with the filenames listed below and "
        "drop them into ThesisTemplate/figure/, everything compiles.")

    _rich_para(doc, [
        ("Target file format: ", {}),
        ("PNG, ≥ 1600 px wide, 200+ DPI, transparent or white background.",
         {"bold": True}),
    ])
    _rich_para(doc, [
        ("Preferred tool: ", {}),
        ("draw.io / diagrams.net ", {"bold": True}),
        ("(save the .drawio source next to the PNG in docs/defense/).",
         {}),
    ])
    _rich_para(doc, [
        ("Font: ", {}),
        ("Inter or Roboto, size ≥ 12 pt", {"bold": True}),
        (" so labels stay legible when scaled to the LaTeX text width.", {}),
    ])

    _hrule(doc)

    # ── Diagram A ──
    _heading(doc, "Diagram A — Channel Transition State Machine (Revision #1)", 2)
    _rich_para(doc, [
        ("Panel comment: ", {"bold": True}),
        ("“Describe the process by which the system determines transitions "
         "among LoRa, LTE, GPRS, WiFi, and offline modes.”",
         {"italic": True}),
    ])
    _rich_para(doc, [
        ("Output filename: ", {"bold": True}),
        ("Fig_5_10_Channel_State_Machine.png", {"name": "Consolas", "size": 10}),
        (" (replaces the placeholder example_gray_box.pdf currently referenced by "
         "\\label{fig:channel_state_machine} in methodology.tex).",
         {}),
    ])
    _rich_para(doc, [
        ("Type: ", {"bold": True}),
        ("finite-state / directed-graph diagram — one node per channel, one edge "
         "per transition trigger. ", {}),
        ("Not", {"italic": True, "bold": True}),
        (" a fall-through flowchart (that already exists as Figure 5.9, "
         "Offline-Buffer Fall-Through Flowchart — do not duplicate it).",
         {}),
    ])

    _heading(doc, "States (nodes)", 3)
    _para(doc,
        "Draw five states in this order, left-to-right, following the priority "
        "chain:")
    for i, s in enumerate([
        "LoRa (SX1278) — highest priority",
        "LTE (A7670E) — cellular preferred",
        "GSM 2G (A7670E) — cellular fallback",
        "WiFi (ESP32) — depot / hotspot",
        "Offline (SPIFFS /tel_buf.log) — buffered",
    ], start=1):
        _rich_bullet(doc, [
            (f"{i}. ", {"bold": True}),
            (s, {}),
        ])
    _para(doc,
        "Use a distinct fill colour per state; keep the palette light. Add a "
        "small subtitle under each node with the underlying hardware (as in "
        "parentheses above) so the reader can map states to the hardware block "
        "diagram.", space_after=8)

    _heading(doc, "Transitions (directed edges)", 3)
    _para(doc,
        "Each edge is labelled with the event that fires it. Use the exact "
        "source in embedded/hmi_phase1/src/main.cpp as the source of truth. "
        "The events are:")
    _table(doc,
        header=["From → To", "Trigger event", "Source-of-truth"],
        rows=[
            ["LoRa → LTE",
             "LoRa ACK timeout (2000 ms) OR LORA_MAX_RETRIES (1) exceeded",
             "LORA_ACK_TIMEOUT_MS, LORA_MAX_RETRIES"],
            ["LTE → GSM 2G",
             "Modem reports RAT_GSM_2G after CGATT / CGACT renegotiation",
             "detectGsmRat(), gsmActiveRat"],
            ["GSM 2G → WiFi",
             "Two consecutive HTTPS POST /telemetry failures on cellular",
             "CH_GSM_NO_SERVER state"],
            ["WiFi → Offline",
             "WiFi.status() != WL_CONNECTED and all cellular attempts failed",
             "bufferTelemetry() call site"],
            ["Offline → any (recovery)",
             "Any live channel returns HTTP 2xx during retry; replay via "
             "flushTelemetryBuffer() at MAX_FLUSH_PER_CALL = 5 records/tick",
             "flushTelemetryBuffer()"],
            ["LoRa (return)",
             "loraGatewayRecentlyOk() — gateway ACK within last 90 s clears "
             "backoff",
             "lastLoraOkMs, loraGatewayBackoffActive()"],
            ["GSM/LTE (return)",
             "AT+CSQ ≥ 10 and AT+CGACT? PDP context active",
             "gsmReconnect()"],
        ],
        col_widths=[1.5, 3.0, 2.0],
    )
    _para(doc, "")
    _rich_para(doc, [
        ("Show recovery edges as ", {}),
        ("dashed arrows", {"bold": True}),
        (" to distinguish them from failure edges (solid). Every state must "
         "have at least one outgoing edge to the next priority level ", {}),
        ("and", {"italic": True}),
        (" a dashed recovery edge back from the offline state (illustrating "
         "that any channel can resume from buffered mode).", {}),
    ], space_after=8)

    _heading(doc, "Annotations", 3)
    _para(doc,
        "Add three small callout boxes anchored to the diagram (not floating "
        "text):")
    _callout(doc,
        "Callout 1 (next to LoRa): “Retry budget: 1 attempt, 2 s ACK window.”")
    _callout(doc,
        "Callout 2 (next to Offline): “Buffer file: /tel_buf.log, ~300 B/record, "
        "~1 MB SPIFFS partition ⇒ ≈ 3 500 records ≈ 4.9 h at 5 s cadence. "
        "32 KB low-space cutoff drops new records to protect the FS.”")
    _callout(doc,
        "Callout 3 (spanning diagram bottom): “Every attempt reuses the same "
        "seq and event_id — replayed records are tagged "
        "channel_used='offline_replay' at the backend so deduplication is safe.”")

    _heading(doc, "Deliverables", 3)
    _bullet(doc, "Fig_5_10_Channel_State_Machine.png in ThesisTemplate/figure/")
    _bullet(doc, "Fig_5_10_Channel_State_Machine.drawio in docs/defense/")
    _bullet(doc,
        "A light-background variant Fig_5_10_Channel_State_Machine_light.png "
        "if the base file is dark-themed (match the pattern of "
        "Fig_5_11_..._light.png).")

    _hrule(doc)

    # ── Diagram B ──
    _heading(doc, "Diagram B — ML Pipeline Redraw (Revision #3)", 2)
    _rich_para(doc, [
        ("Panel comment: ", {"bold": True}),
        ("“Revise the diagram (Figure 5.1) with corrected arrows, greater "
         "detail, and a general representation in place of Python.”",
         {"italic": True}),
    ])
    _rich_para(doc, [
        ("⚠  Ambiguity to resolve first: ", {"bold": True, "color": (0xB4, 0x4A, 0x00)}),
        ("the panel wrote ", {}),
        ("“Figure 5.1”", {"italic": True}),
        (" but the current LaTeX numbering has:", {}),
    ])
    _bullet(doc, "Figure 5.1 = concep_diagram.png (Conceptual Framework — "
                 "no Python content)")
    _bullet(doc, "Figure 5.11 = Fig_5_11_Anomaly_Pipeline_light.png "
                 "(ML pipeline — has Python/library icons)")
    _bullet(doc, "Sub-figure inside 5.11 group = thesis_defense_pipeline_v2.png "
                 "(also ML pipeline)")
    _rich_para(doc, [
        ("The ", {}),
        ("“in place of Python”", {"italic": True}),
        (" cue points at ", {}),
        ("Figure 5.11", {"bold": True}),
        (" (or its companion thesis_defense_pipeline_v2.png). Confirm with the "
         "panel chair before starting; assume 5.11 by default.", {}),
    ])
    _rich_para(doc, [
        ("Output filename: ", {"bold": True}),
        ("Fig_5_11_Anomaly_Pipeline_v3.png ", {"name": "Consolas", "size": 10}),
        ("(keeps the naming convention; just bump the version suffix). Add a "
         "_light.png variant too.", {}),
    ])

    _heading(doc, "What to fix", 3)
    _rich_para(doc, [
        ("1.  ", {"bold": True}),
        ("“Corrected arrows.” ", {"bold": True}),
        ("The current diagram has ambiguous edge directions between the "
         "primary inference plane and the audit/override inference plane. "
         "Redraw so every arrow has a single direction and a labelled event:",
         {}),
    ])
    _bullet(doc, "IF pass → contextual gate  (label: if_score ≥ τ + δ)")
    _bullet(doc, "MP override → contextual gate  (label: mp_score ≥ 4.0 ∧ "
                 "Δfuel < 0  OR  mp_score ≥ 2.0 ∧ Δfuel ≤ −2.0)")
    _bullet(doc, "Contextual gate → persistence  (label: pass)")
    _bullet(doc, "Contextual gate → audit-only sink  (label: suppressed)")
    _bullet(doc, "Persistence → alerts table  (label: combined_score > 0.65 "
                 "∧ dedup ok)")

    _rich_para(doc, [
        ("2.  ", {"bold": True}),
        ("“Greater detail.” ", {"bold": True}),
        ("Add the four suppressors as labelled boxes inside the contextual "
         "gate (currently collapsed into one box):", {}),
    ])
    _bullet(doc, "parkedRefuel — rejects fill-up (fuel Δ > +10% while parked)")
    _bullet(doc, "postRefuel — 5-min settle window after any +5% fuel jump")
    _bullet(doc, "loadSpikeSuppress — engine load ≥ 55% AND load rise ≥ 25% "
                 "with small negative fuel Δ")
    _bullet(doc, "coldStartSuppress — sender-priming interval after ignition")

    _rich_para(doc, [
        ("3.  ", {"bold": True}),
        ("“General representation in place of Python.” ", {"bold": True}),
        ("Remove any Python-specific icons/labels currently in the PNG and "
         "replace with generic algorithm names:", {}),
    ])
    _table(doc,
        header=["Remove (Python-specific)", "Replace with (general)"],
        rows=[
            ["sklearn.IsolationForest",
             "Isolation Forest (one-class ensemble)"],
            ["stumpy.matrix_profile",
             "Matrix Profile (subsequence discord)"],
            ["pandas.DataFrame / numpy.ndarray icons",
             "Feature-vector row (33 dim)"],
            ["pickle model file icon",
             "Serialised model artefact"],
            ["any Python code snippet, sklearn / stumpy / numpy logos, snake "
             "icons",
             "plain algorithm names + arrows"],
        ],
        col_widths=[3.2, 3.3],
    )
    _para(doc, "")
    _rich_para(doc, [
        ("Also drop any file-extension hints like .pkl, .py, .csv in favour "
         "of the conceptual data name (e.g. ", {}),
        ("“trained model”, “feature row”", {"italic": True}),
        (").", {}),
    ], space_after=8)

    _heading(doc, "Deliverables", 3)
    _bullet(doc, "Fig_5_11_Anomaly_Pipeline_v3.png and _light.png in "
                 "ThesisTemplate/figure/")
    _bullet(doc, "Fig_5_11_Anomaly_Pipeline_v3.drawio in docs/defense/")
    _bullet(doc, "Confirm the LaTeX \\includegraphics{...} in methodology.tex "
                 "around line 1071 points at the new filename (Sam will "
                 "handle the swap once the PNG lands).")

    _hrule(doc)

    # ── Wrap-up ──
    _heading(doc, "Once your PNGs are dropped in ThesisTemplate/figure/", 2)
    _para(doc, "Ping Sam. He will:")
    _bullet(doc,
        "Swap the placeholder example_gray_box.pdf reference to your PNG in "
        "methodology.tex (both figures).")
    _bullet(doc,
        "Re-flow the caption / paragraph to match the final visual (the "
        "placeholder captions and surrounding prose are already written but "
        "may need one-line edits once we see the actual layout).")
    _bullet(doc,
        "Rebuild the PDF and verify figure numbering renumbered correctly "
        "(Fig 5.10 is inserted between the existing 5.9 and 5.11, so "
        "cross-references downstream might shift by one).")
    _para(doc,
        "Reach out on Messenger if any transition trigger in the state-machine "
        "table is unclear — the firmware source is the authority, but the "
        "phrasing above should be enough to draw the diagram without reading "
        "C++.", italic=True)

    doc.save(OUT)
    print(f"[docx] wrote {OUT}")


if __name__ == "__main__":
    build()
